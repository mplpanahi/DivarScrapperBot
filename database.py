import os
import sqlite3
import json
import csv
from datetime import datetime
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class Database:
    """SQLite Database manager for Divar scraped ads."""

    def __init__(self, db_path: str = "data/divar_scraper.db"):
        self.db_path = db_path
        os.makedirs(os.path.dirname(self.db_path), exist_ok=True)
        self._init_db()

    def _get_connection(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def _init_db(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                CREATE TABLE IF NOT EXISTS ads (
                    token TEXT PRIMARY KEY,
                    title TEXT,
                    price TEXT,
                    phone_number TEXT,
                    city TEXT,
                    city_persian TEXT,
                    district TEXT,
                    category TEXT,
                    description TEXT,
                    image_url TEXT,
                    images_json TEXT,
                    attributes_json TEXT,
                    url TEXT,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_ads_phone ON ads(phone_number)")
            cursor.execute("CREATE INDEX IF NOT EXISTS idx_ads_city ON ads(city)")
            conn.commit()

    def has_ad(self, token: str) -> bool:
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT 1 FROM ads WHERE token = ?", (token,))
            return cursor.fetchone() is not None

    def save_ad(self, ad: Dict[str, Any]) -> bool:
        """
        Saves or updates an ad.
        Returns True if inserted/updated.
        """
        token = ad.get("token")
        if not token:
            return False

        images_json = json.dumps(ad.get("images", []), ensure_ascii=False)
        attributes_json = json.dumps(ad.get("attributes", {}), ensure_ascii=False)

        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO ads (
                    token, title, price, phone_number, city, city_persian,
                    district, category, description, image_url, images_json,
                    attributes_json, url, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(token) DO UPDATE SET
                    title = excluded.title,
                    price = COALESCE(NULLIF(excluded.price, ''), ads.price),
                    phone_number = COALESCE(NULLIF(excluded.phone_number, ''), ads.phone_number),
                    district = COALESCE(NULLIF(excluded.district, ''), ads.district),
                    description = COALESCE(NULLIF(excluded.description, ''), ads.description),
                    attributes_json = COALESCE(NULLIF(excluded.attributes_json, '{}'), ads.attributes_json),
                    images_json = COALESCE(NULLIF(excluded.images_json, '[]'), ads.images_json)
            """, (
                token,
                ad.get("title", ""),
                ad.get("price", ""),
                ad.get("phone_number", ""),
                ad.get("city", ""),
                ad.get("city_persian", ""),
                ad.get("district", ""),
                ad.get("category", ""),
                ad.get("description", ""),
                ad.get("image_url", ""),
                images_json,
                attributes_json,
                ad.get("url", f"https://divar.ir/v/{token}"),
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ))
            conn.commit()
            return True

    def get_ads(self, limit: int = 100, offset: int = 0, with_phone_only: bool = False) -> List[Dict[str, Any]]:
        with self._get_connection() as conn:
            cursor = conn.cursor()
            query = "SELECT * FROM ads"
            params = []
            if with_phone_only:
                query += " WHERE phone_number IS NOT NULL AND phone_number != '' AND phone_number NOT LIKE '%چت%'"
            query += " ORDER BY created_at DESC LIMIT ? OFFSET ?"
            params.extend([limit, offset])

            cursor.execute(query, params)
            rows = cursor.fetchall()
            result = []
            for r in rows:
                d = dict(r)
                try:
                    d["attributes"] = json.loads(d.get("attributes_json") or "{}")
                except Exception:
                    d["attributes"] = {}
                result.append(d)
            return result

    def get_counts(self) -> Dict[str, int]:
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*) FROM ads")
            total = cursor.fetchone()[0]

            cursor.execute("SELECT COUNT(*) FROM ads WHERE phone_number IS NOT NULL AND phone_number != '' AND phone_number NOT LIKE '%چت%'")
            with_phone = cursor.fetchone()[0]

            return {
                "total_ads": total,
                "with_phone": with_phone,
                "without_phone": total - with_phone
            }

    def clear_ads(self):
        with self._get_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("DELETE FROM ads")
            conn.commit()

    def export_to_excel(self, output_path: str = "output/divar_ads.xlsx") -> str:
        """Exports all ads to a beautifully formatted Excel file with Persian font and borders."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        ads = self.get_ads(limit=100000)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "آگهی‌های دیوار"
        ws.sheet_view.rightToLeft = True

        headers = [
            "ردیف", "عنوان آگهی", "شماره تماس", "قیمت", "شهر", 
            "محله", "دسته‌بندی", "لینک آگهی", "تاریخ ثبت در ربات", "توضیحات"
        ]
        ws.append(headers)

        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
        cell_font = Font(name="Tahoma", size=10)
        phone_font = Font(name="Tahoma", size=11, bold=True, color="047857")
        center_align = Alignment(horizontal="center", vertical="center")
        right_align = Alignment(horizontal="right", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='E5E7EB'),
            right=Side(style='thin', color='E5E7EB'),
            top=Side(style='thin', color='E5E7EB'),
            bottom=Side(style='thin', color='E5E7EB')
        )

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for i, ad in enumerate(ads, start=1):
            phone = ad.get("phone_number", "")
            row_data = [
                i,
                ad.get("title", ""),
                phone if phone else "در دسترس نیست",
                ad.get("price", "توافقی"),
                ad.get("city_persian") or ad.get("city", ""),
                ad.get("district", ""),
                ad.get("category", ""),
                ad.get("url", ""),
                ad.get("created_at", ""),
                (ad.get("description", "") or "")[:250]
            ]
            ws.append(row_data)
            row_num = i + 1
            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=col_num)
                c.font = cell_font
                c.border = thin_border
                if col_num in (1, 5, 6, 7, 9):
                    c.alignment = center_align
                elif col_num == 3:
                    c.font = phone_font
                    c.alignment = center_align
                else:
                    c.alignment = right_align

        # Adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = min(len(val), 40)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(output_path)
        return output_path

    def export_to_csv(self, output_path: str = "output/divar_ads.csv") -> str:
        """Exports all ads to a UTF-8 with BOM CSV file for Excel compatibility."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        ads = self.get_ads(limit=100000)

        with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f)
            writer.writerow([
                "ردیف", "عنوان", "شماره تماس", "قیمت", "شهر", "محله", 
                "دسته‌بندی", "لینک آگهی", "تاریخ ثبت", "توضیحات"
            ])
            for i, ad in enumerate(ads, start=1):
                writer.writerow([
                    i,
                    ad.get("title", ""),
                    ad.get("phone_number", ""),
                    ad.get("price", ""),
                    ad.get("city_persian") or ad.get("city", ""),
                    ad.get("district", ""),
                    ad.get("category", ""),
                    ad.get("url", ""),
                    ad.get("created_at", ""),
                    (ad.get("description", "") or "")[:200]
                ])
        return output_path
